"""F3 structured-export regression tests."""

# pyright: reportMissingTypeStubs=false, reportUnknownArgumentType=false, reportUnknownMemberType=false, reportUnknownVariableType=false

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import load_workbook

from qf_platform.application import (
    HedgeWorkbenchConfig,
    HedgeWorkbenchRequest,
    HestonCalibrationDraft,
    M2WorkbenchAnalysis,
    M2WorkbenchConfig,
    M2WorkbenchRequest,
    MarketWorkbenchConfig,
    WorkbenchValuationMethod,
    canonical_black_scholes_draft,
    canonical_m4_market_workbench,
    canonical_m8_performance_reference,
    compose_black_scholes_study,
    make_heston_calibration_request,
    make_ui5_reference_validation_request,
    run_hedge_workbench,
    run_heston_calibration,
    run_m2_workbench,
    run_ui5_validation,
)
from qf_platform.reporting import (
    export_report,
    hedging_report,
    heston_calibration_report,
    market_iv_report,
    validation_model_risk_report,
    valuation_greeks_report,
)
from qf_platform.sensitivity import (
    BlackScholesSensitivity,
    FiniteDifferenceBlackScholesSensitivity,
)


def _m2_analysis() -> tuple[M2WorkbenchRequest, M2WorkbenchAnalysis]:
    composition = compose_black_scholes_study(canonical_black_scholes_draft())
    request = M2WorkbenchRequest(
        composition=composition,
        config=M2WorkbenchConfig(
            valuation_method=WorkbenchValuationMethod.ANALYTIC,
            crr_steps=200,
            monte_carlo_paths=2_000,
            monte_carlo_seed=1729,
            selected_greek=BlackScholesSensitivity.DELTA,
            finite_difference_method=FiniteDifferenceBlackScholesSensitivity(
                spot_bump=0.1,
                volatility_bump=0.001,
                rate_bump=0.0001,
                theta_day_bump=1,
            ),
        ),
    )
    return request, run_m2_workbench(request)


def test_concrete_adapters_preserve_authoritative_values() -> None:
    m2_request, m2_analysis = _m2_analysis()
    valuation = valuation_greeks_report(m2_request, m2_analysis)
    valuations = next(
        table for table in valuation.tables if table.name == "valuations"
    )
    analytic = next(
        run for run in m2_analysis.valuations if run.method.value == "analytic"
    )
    assert analytic.result is not None
    assert valuations.rows[0][3] == analytic.result.present_value

    composition = compose_black_scholes_study(canonical_black_scholes_draft())
    hedge_request = HedgeWorkbenchRequest(
        composition=composition,
        config=HedgeWorkbenchConfig(0.20, 0.20, 7, 1729, 4, 0.001),
    )
    hedge_analysis = run_hedge_workbench(hedge_request)
    hedge = hedging_report(hedge_request, hedge_analysis)
    replicates = next(
        table for table in hedge.tables if table.name == "replicates"
    )
    assert (
        replicates.rows[0][1]
        == hedge_analysis.selected_replicates[0].replication_error
    )

    market_analysis = canonical_m4_market_workbench()
    market = market_iv_report(MarketWorkbenchConfig(), market_analysis)
    observations = next(
        table for table in market.tables if table.name == "observations"
    )
    first_outcome = market_analysis.outcomes[0]
    assert first_outcome.normalized is not None
    assert observations.rows[0][6] == first_outcome.normalized.target_price

    calibration_request = make_heston_calibration_request(
        "thin",
        HestonCalibrationDraft(),
    )
    calibration_analysis = run_heston_calibration(calibration_request)
    calibration = heston_calibration_report(
        calibration_request,
        calibration_analysis,
    )
    runs = next(
        table for table in calibration.tables if table.name == "calibration_runs"
    )
    assert (
        runs.rows[0][11]
        == calibration_analysis.runs[0].result.objective_value
    )


def test_flagship_validation_exports_are_auditable(tmp_path: Path) -> None:
    request = make_ui5_reference_validation_request()
    analysis = run_ui5_validation(request)
    performance = canonical_m8_performance_reference()
    report = validation_model_risk_report(
        request,
        analysis,
        performance=performance,
    )

    names = tuple(table.name for table in report.tables)
    assert names == (
        "summary",
        "training_metrics",
        "heldout_metrics",
        "training_contracts",
        "heldout_contracts",
        "residuals",
        "calibrated_parameters",
        "calibration_starts",
        "stability",
        "conditioning",
        "assumptions",
        "provenance",
        "performance",
    )
    heldout = next(
        table for table in report.tables if table.name == "heldout_metrics"
    )
    assert (
        heldout.rows[0][4]
        == analysis.evidence.black_scholes_evaluation_metrics.root_mean_square_error
    )
    assert (
        heldout.rows[1][4]
        == analysis.evidence.heston_evaluation_metrics.root_mean_square_error
    )

    paths = export_report(report, tmp_path)
    assert paths.xlsx.is_file()
    assert paths.json.is_file()
    assert all(path.is_file() for path in paths.csv)

    workbook = load_workbook(paths.xlsx, data_only=False, read_only=True)
    assert "training_metrics" in workbook.sheetnames
    assert "heldout_metrics" in workbook.sheetnames
    assert "conditioning" in workbook.sheetnames
    assert "provenance" in workbook.sheetnames
    assert workbook["heldout_metrics"]["E2"].value == heldout.rows[0][4]

    payload = json.loads(paths.json.read_text(encoding="utf-8"))
    assert payload["schema"] == "qf-platform-tabular-report-v1"
    tables = {table["name"]: table for table in payload["tables"]}
    assert tables["heldout_metrics"]["rows"][0]["rmse"] == heldout.rows[0][4]
    assert (
        tables["performance"]["rows"][2]["speedup_x"]
        == performance.workloads[2].speedup_x
    )

    metadata_csv = next(
        path for path in paths.csv if path.name == "_metadata.csv"
    )
    with metadata_csv.open(encoding="utf-8", newline="") as handle:
        rows = list(csv.reader(handle))
    assert ["design", "same-date cross-sectional holdout", "", ""] in rows
