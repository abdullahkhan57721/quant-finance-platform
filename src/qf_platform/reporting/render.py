"""XLSX, CSV, and JSON renderers for explicit tabular report values."""

# pyright: reportMissingTypeStubs=false, reportUnknownArgumentType=false, reportUnknownMemberType=false, reportUnknownVariableType=false

from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from pathlib import Path

from .model import ReportField, ReportTable, ReportValue, TabularReport


@dataclass(frozen=True, slots=True)
class ExportedReportPaths:
    """Filesystem paths produced by one all-format tabular report export."""

    xlsx: Path
    json: Path
    csv: tuple[Path, ...]


def _field_record(field: ReportField) -> dict[str, ReportValue]:
    return {
        "key": field.key,
        "value": field.value,
        "unit": field.unit,
        "description": field.description,
    }


def _table_record(table: ReportTable) -> dict[str, object]:
    return {
        "name": table.name,
        "columns": [
            {
                "key": column.key,
                "label": column.label,
                "unit": column.unit,
                "description": column.description,
            }
            for column in table.columns
        ],
        "rows": [
            {\n                column.key: value\n                for column, value in zip(table.columns, row, strict=True)\n            }
            for row in table.rows
        ],
    }


def write_json_report(report: TabularReport, path: str | Path) -> Path:
    """Write machine-readable values without display rounding."""

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "schema": "qf-platform-tabular-report-v1",
        "report_id": report.report_id,
        "title": report.title,
        "metadata": [_field_record(field) for field in report.metadata],
        "tables": [_table_record(table) for table in report.tables],
    }
    destination.write_text(
        json.dumps(payload, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return destination


def write_csv_report(report: TabularReport, directory: str | Path) -> tuple[Path, ...]:
    """Write one CSV per meaningful rectangular table plus report metadata."""

    destination = Path(directory)
    destination.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []

    metadata_path = destination / "_metadata.csv"
    with metadata_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(("key", "value", "unit", "description"))
        for field in report.metadata:
            writer.writerow((field.key, field.value, field.unit, field.description))
    written.append(metadata_path)

    for table in report.tables:
        table_path = destination / f"{table.name}.csv"
        with table_path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(tuple(column.display_label for column in table.columns))
            writer.writerows(table.rows)
        written.append(table_path)
    return tuple(written)


def _autosize(worksheet: object) -> None:
    for column_cells in worksheet.columns:  # type: ignore[attr-defined]
        maximum = max(
            (\n                len(str(cell.value)) if cell.value is not None else 0\n                for cell in column_cells\n            ),
            default=0,
        )
        letter = column_cells[0].column_letter
        worksheet.column_dimensions[letter].width = min(max(maximum + 2, 10), 48)  # type: ignore[attr-defined]


def write_xlsx_report(report: TabularReport, path: str | Path) -> Path:
    """Write an analyst-friendly workbook over raw report values, not formulas."""

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:
        raise RuntimeError(
            "XLSX export requires the optional reporting dependencies; "
            "install quant-finance-platform[reporting]"
        ) from exc

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    metadata_sheet = workbook.active
    if metadata_sheet is None:
        raise RuntimeError("openpyxl workbook did not create an active worksheet")
    metadata_sheet.title = "_metadata"
    metadata_sheet.append(("report_id", report.report_id, "", "stable report identity"))
    metadata_sheet.append(("title", report.title, "", "human-readable report title"))
    metadata_sheet.append(("key", "value", "unit", "description"))
    for cell in metadata_sheet[3]:
        cell.font = Font(bold=True)
    for field in report.metadata:
        metadata_sheet.append((field.key, field.value, field.unit, field.description))
    metadata_sheet.freeze_panes = "A4"
    _autosize(metadata_sheet)

    for table in report.tables:
        sheet = workbook.create_sheet(table.name)
        sheet.append(tuple(column.display_label for column in table.columns))
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in table.rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        _autosize(sheet)

    workbook.save(destination)
    return destination


def export_report(
    report: TabularReport,
    directory: str | Path,
    *,
    stem: str | None = None,
) -> ExportedReportPaths:
    """Render one tabular report to XLSX, JSON, and a CSV directory."""

    destination = Path(directory)
    destination.mkdir(parents=True, exist_ok=True)
    resolved_stem = stem or report.report_id
    xlsx = write_xlsx_report(report, destination / f"{resolved_stem}.xlsx")
    json_path = write_json_report(report, destination / f"{resolved_stem}.json")
    csv_paths = write_csv_report(report, destination / f"{resolved_stem}_csv")
    return ExportedReportPaths(xlsx=xlsx, json=json_path, csv=csv_paths)
